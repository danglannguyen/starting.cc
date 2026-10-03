"""Tests for ats.parsing."""

import io
import zipfile

import docx
import pytest

from ats import parsing
from ats.config import MAX_DOCUMENT_CHARS
from ats.parsing import (
    TEXT_BOX_WARNING,
    ParsingError,
    normalize_text,
    read_cv_docx,
    read_jd_file,
    read_jd_text,
)
from tests.helpers import FIXTURES_DIR, add_text_box, build_pdf, docx_to_bytes

# --- normalize_text ---------------------------------------------------------


def test_normalize_lowercases_and_collapses_whitespace() -> None:
    assert normalize_text("  Lean   Six\tSigma\n\nSQL  ") == "lean six sigma sql"


def test_normalize_treats_non_breaking_space_as_space() -> None:
    assert normalize_text("Power BI") == "power bi"


def test_normalize_expands_pdf_ligatures() -> None:
    # "ﬁ" is the single "fi" ligature character often found in PDFs.
    assert normalize_text("Certiﬁed") == "certified"


# --- pasted text and .txt ---------------------------------------------------


def test_pasted_text_keeps_original_and_normalizes_copy() -> None:
    result = read_jd_text("  Senior Analyst\nSQL  ")
    assert result.original_text == "Senior Analyst\nSQL"
    assert result.normalized_text == "senior analyst sql"
    assert result.warnings == []


def test_empty_pasted_text_raises() -> None:
    with pytest.raises(ParsingError, match="No readable text"):
        read_jd_text("   \n  ")


def test_text_over_the_length_limit_raises() -> None:
    too_long = "a" * (MAX_DOCUMENT_CHARS + 1)
    with pytest.raises(ParsingError, match="too long"):
        read_jd_text(too_long)


def test_text_at_the_length_limit_is_accepted() -> None:
    result = read_jd_text("a" * MAX_DOCUMENT_CHARS)
    assert len(result.original_text) == MAX_DOCUMENT_CHARS


def test_txt_fixture_file_is_read() -> None:
    file_bytes = (FIXTURES_DIR / "sample_jd.txt").read_bytes()
    result = read_jd_file(file_bytes, "sample_jd.txt")
    assert "lean six sigma green belt (lss)" in result.normalized_text
    assert result.original_text.startswith("Job Title:")


def test_txt_with_windows_bom_is_read_without_the_marker() -> None:
    result = read_jd_file("﻿SQL required".encode(), "jd.txt")
    assert result.original_text == "SQL required"


def test_txt_that_is_not_utf8_raises_friendly_error() -> None:
    latin1_bytes = "Café manager".encode("latin-1")
    with pytest.raises(ParsingError, match="not UTF-8"):
        read_jd_file(latin1_bytes, "jd.txt")


# --- file type dispatch -----------------------------------------------------


def test_unsupported_extension_raises() -> None:
    with pytest.raises(ParsingError, match="Unsupported"):
        read_jd_file(b"anything", "jd.doc")


def test_extension_check_ignores_case() -> None:
    result = read_jd_file(build_pdf(["Python developer"]), "JD.PDF")
    assert "python developer" in result.normalized_text


# --- PDF --------------------------------------------------------------------


def test_pdf_text_is_extracted() -> None:
    pdf_bytes = build_pdf(["Required: SQL and Power BI", "Nice to have: Python"])
    result = read_jd_file(pdf_bytes, "jd.pdf")
    assert "required: sql and power bi" in result.normalized_text
    assert "nice to have: python" in result.normalized_text


def test_pdf_without_text_layer_raises_scanned_hint() -> None:
    with pytest.raises(ParsingError, match="scanned image"):
        read_jd_file(build_pdf([]), "jd.pdf")


def test_pdf_over_the_page_limit_raises(monkeypatch: pytest.MonkeyPatch) -> None:
    # Lower the limit instead of building a 21-page PDF.
    monkeypatch.setattr(parsing, "MAX_PDF_PAGES", 0)
    with pytest.raises(ParsingError, match="limit is 0"):
        read_jd_file(build_pdf(["Python developer"]), "jd.pdf")


def test_damaged_pdf_raises_friendly_error() -> None:
    with pytest.raises(ParsingError, match="could not be opened"):
        read_jd_file(b"this is not a pdf", "jd.pdf")


# --- .docx ------------------------------------------------------------------


def test_docx_paragraphs_are_read_and_case_kept_for_display() -> None:
    document = docx.Document()
    document.add_paragraph("Jane Example")
    document.add_paragraph("Process Improvement Analyst")
    result = read_cv_docx(docx_to_bytes(document))
    assert result.original_text == "Jane Example\nProcess Improvement Analyst"
    assert result.normalized_text == "jane example process improvement analyst"
    assert result.warnings == []


def test_docx_table_text_is_read() -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "SQL, Power BI"
    table.cell(1, 0).text = "Certifications"
    table.cell(1, 1).text = "Lean Six Sigma Green Belt"
    result = read_cv_docx(docx_to_bytes(document))
    assert "Skills | SQL, Power BI" in result.original_text
    assert "Certifications | Lean Six Sigma Green Belt" in result.original_text


def test_docx_nested_table_text_is_read() -> None:
    document = docx.Document()
    outer_cell = document.add_table(rows=1, cols=1).cell(0, 0)
    outer_cell.text = "Tools"
    outer_cell.add_table(rows=1, cols=1).cell(0, 0).text = "Minitab"
    result = read_cv_docx(docx_to_bytes(document))
    assert "minitab" in result.normalized_text


def test_docx_merged_cell_text_appears_once() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Experience"
    table.cell(0, 2).text = "2019-2024"
    result = read_cv_docx(docx_to_bytes(document))
    assert result.original_text == "Experience | 2019-2024"


def test_docx_header_and_footer_are_read() -> None:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Example | jane@example.com"
    document.sections[0].footer.paragraphs[0].text = "Page footer note"
    document.add_paragraph("Body text")
    result = read_cv_docx(docx_to_bytes(document))
    assert result.original_text == (
        "Jane Example | jane@example.com\nBody text\nPage footer note"
    )


def test_docx_header_shared_by_two_sections_is_read_once() -> None:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Shared header"
    document.add_paragraph("Section one")
    document.add_section()  # new sections reuse ("link to") the previous header
    document.add_paragraph("Section two")
    result = read_cv_docx(docx_to_bytes(document))
    assert result.original_text.count("Shared header") == 1


def test_docx_text_box_is_flagged_and_not_read() -> None:
    document = docx.Document()
    document.add_paragraph("Visible summary")
    add_text_box(document)
    result = read_cv_docx(docx_to_bytes(document))
    assert result.warnings == [TEXT_BOX_WARNING]
    assert "kubernetes" not in result.normalized_text
    assert "visible summary" in result.normalized_text


def test_docx_with_only_text_box_text_raises() -> None:
    document = docx.Document()
    add_text_box(document)
    with pytest.raises(ParsingError, match="inside text boxes"):
        read_cv_docx(docx_to_bytes(document))


def test_empty_docx_raises() -> None:
    with pytest.raises(ParsingError, match="No readable text"):
        read_cv_docx(docx_to_bytes(docx.Document()))


def test_non_docx_bytes_raise_friendly_error() -> None:
    with pytest.raises(ParsingError, match=r"could not be opened as a Word"):
        read_cv_docx(b"this is not a docx file")


def test_docx_with_damaged_xml_raises_friendly_error() -> None:
    # A valid zip whose main document XML is cut in half.
    document = docx.Document()
    document.add_paragraph("Synthetic text")
    damaged = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(docx_to_bytes(document))) as source,
        zipfile.ZipFile(damaged, "w") as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = data[: len(data) // 2]
            target.writestr(item, data)
    with pytest.raises(ParsingError, match=r"could not be opened as a Word"):
        read_cv_docx(damaged.getvalue())
